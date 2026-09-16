"""주문 접수 — 폼 응답을 job.json 으로.

스마트스토어 주문서에는 파일을 첨부할 수 없다. 그래서 결제 후 폼 링크를
안내해 사진과 문구를 따로 받는 동선이 된다. 그 응답을 손으로 job.json 에
옮겨 적으면 주문 1건마다 15~20개 필드를 타이핑하게 되고, 오타는 그대로
납품물에 나간다. 주문제작에서 사람 시간이 가장 많이 새는 지점이다.

폼 컬럼 제목은 판매자가 자유롭게 쓰므로 자동 추측은 믿을 게 못 된다.
매핑 파일을 한 번 만들어 두고 계속 재사용한다 — 폼을 고치지 않는 한
다음 주문부터는 명령 한 줄이다.
"""
from __future__ import annotations
import csv, json, re
from dataclasses import dataclass, field
from pathlib import Path

from engine.spec import Template

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}
VIDEO_EXT = {".mp4", ".mov", ".m4v"}


# ── 표 읽기 ───────────────────────────────────────────────────────────────
def read_table(path: str | Path) -> tuple[list[str], list[dict[str, str]]]:
    """xlsx / csv → (컬럼 목록, 행 목록). 네이버폼은 xlsx 로 내려받는다."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("xlsx 를 읽으려면 openpyxl 이 필요합니다: "
                               "pip install openpyxl")
        ws = load_workbook(p, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        cols = [str(c).strip() if c is not None else "" for c in rows[0]]
        out = []
        for r in rows[1:]:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            out.append({cols[i]: ("" if v is None else str(v).strip())
                        for i, v in enumerate(r) if i < len(cols)})
        return cols, out

    # csv — 네이버폼에서 csv 로 내보냈거나 시트에서 복사한 경우
    text = p.read_text(encoding="utf-8-sig")
    rdr = csv.DictReader(text.splitlines())
    cols = [c.strip() for c in (rdr.fieldnames or [])]
    rows = [{(k or "").strip(): (v or "").strip() for k, v in r.items()}
            for r in rdr]
    rows = [r for r in rows if any(r.values())]
    return cols, rows


# ── 매핑 ──────────────────────────────────────────────────────────────────
def _norm(s: str) -> str:
    return re.sub(r"[\s()\[\]·:,./-]+", "", str(s)).lower()


def suggest_mapping(cols: list[str], tpl: Template) -> dict[str, str]:
    """컬럼 제목과 슬롯을 best-effort 로 짝지어 본다.

    어디까지나 초안이다. 사람이 확인하라고 만든 것이지 믿고 쓰라는 게 아니다.
    """
    out: dict[str, str] = {}
    taken: set[str] = set()
    for spec in tpl.inputs:
        if spec.type != "text":
            continue
        cands = [_norm(spec.id), _norm(spec.guide or "")]
        best, score = None, 0
        for c in cols:
            if not c or c in taken:
                continue
            nc = _norm(c)
            for cand in cands:
                if not cand:
                    continue
                if nc == cand:
                    s = 3
                elif cand in nc or nc in cand:
                    s = 2
                else:
                    s = 0
                if s > score:
                    best, score = c, s
        if best and score >= 2:
            out[spec.id] = best
            taken.add(best)
        else:
            out[spec.id] = ""          # 사람이 채운다
    return out


def make_mapping(tpl: Template, cols: list[str]) -> dict:
    return {
        "template": tpl.id,
        "_컬럼목록": cols,             # 편집할 때 보라고 남겨둔다
        "order_id": {"column": "", "prefix": "ORDER"},
        "photos": {
            "root": "storage/uploads",
            "folder_column": "",       # 비우면 order_id 로 폴더를 찾는다
            "order": "filename",       # filename | mtime
        },
        "text": suggest_mapping(cols, tpl),
    }


# ── job 생성 ──────────────────────────────────────────────────────────────
@dataclass
class Built:
    order_id: str
    path: Path | None
    photos: int
    problems: list[str] = field(default_factory=list)


def _clean(v: str, spec) -> str:
    """폼 입력을 그대로 쓰면 안 되는 것들만 정리한다."""
    v = (v or "").strip()
    v = v.replace("\r\n", "\n").replace("\r", "\n")
    v = v.replace("\\n", "\n")                 # 고객이 문자 그대로 \n 을 쓴 경우
    v = re.sub(r"\n{3,}", "\n\n", v)
    v = re.sub(r"[ \t]+\n", "\n", v)
    return v


def collect_photos(folder: Path, order: str = "filename") -> list[Path]:
    if not folder.is_dir():
        return []
    files = [p for p in folder.iterdir()
             if p.is_file() and p.suffix.lower() in IMAGE_EXT | VIDEO_EXT]
    if order == "mtime":
        files.sort(key=lambda p: p.stat().st_mtime)
    else:
        # 고객이 1_.jpg, 2_.jpg 로 번호를 매기면 그 순서가 된다.
        # 숫자를 숫자로 비교해야 10 이 2 뒤에 온다.
        files.sort(key=lambda p: [int(t) if t.isdigit() else t.lower()
                                  for t in re.split(r"(\d+)", p.name)])
    return files


def build_jobs(rows: list[dict], tpl: Template, mapping: dict,
               out_dir: Path, *, aspect: str | None = None,
               dry_run: bool = False) -> list[Built]:
    out_dir.mkdir(parents=True, exist_ok=True)
    oid_cfg = mapping.get("order_id", {})
    ph_cfg = mapping.get("photos", {})
    root = Path(ph_cfg.get("root") or "storage/uploads")
    text_map = mapping.get("text", {})
    specs = {s.id: s for s in tpl.inputs}
    media_slots = [s.id for s in tpl.inputs if s.type in ("image", "video")]
    required_media = [s.id for s in tpl.inputs
                      if s.type in ("image", "video") and s.required]

    built: list[Built] = []
    for i, row in enumerate(rows, 1):
        problems: list[str] = []

        oid = (row.get(oid_cfg.get("column", ""), "") or "").strip()
        if not oid:
            oid = f"{oid_cfg.get('prefix', 'ORDER')}-{i:03d}"
        oid = re.sub(r"[^\w.-]+", "_", oid)

        inputs: dict[str, str] = {}
        for slot, col in text_map.items():
            if not col:
                continue
            spec = specs.get(slot)
            if spec is None:
                problems.append(f"매핑에 없는 슬롯: {slot}")
                continue
            val = _clean(row.get(col, ""), spec)
            if val:
                inputs[slot] = val
            elif spec.required and not spec.default:
                problems.append(f"필수 문구 비어 있음: {slot} (컬럼 '{col}')")

        unmapped = [s.id for s in tpl.inputs
                    if s.type == "text" and s.required and not s.default
                    and not text_map.get(s.id)]
        for s in unmapped:
            problems.append(f"필수 문구가 매핑되지 않음: {s}")

        folder_col = ph_cfg.get("folder_column") or ""
        folder_name = (row.get(folder_col, "") or "").strip() if folder_col else ""
        folder = root / (folder_name or oid)
        pics = collect_photos(folder, ph_cfg.get("order", "filename"))
        if not pics:
            problems.append(f"사진 폴더가 비었거나 없음: {folder}")
        for slot, pic in zip(media_slots, pics):
            inputs[slot] = str(pic)
        if len(pics) < len(required_media):
            problems.append(f"필수 사진 {len(required_media)}장 중 {len(pics)}장만 있음")
        elif len(pics) > len(media_slots):
            problems.append(f"사진 {len(pics)}장 중 앞 {len(media_slots)}장만 사용")

        job = {"template": tpl.id, "order_id": oid, "preview": True,
               "inputs": inputs}
        if aspect:
            job["aspect"] = aspect

        path = out_dir / f"{oid}.json"
        if not dry_run:
            path.write_text(json.dumps(job, ensure_ascii=False, indent=2),
                            encoding="utf-8")
        built.append(Built(oid, None if dry_run else path, len(pics), problems))
    return built


def report(built: list[Built]) -> str:
    lines = []
    bad = 0
    for b in built:
        mark = "✖" if b.problems else "·"
        lines.append(f"{mark} {b.order_id:<16} 사진 {b.photos:>2}장"
                     + (f"  → {b.path}" if b.path else "  (dry-run)"))
        for p in b.problems:
            lines.append(f"{'':>2}  └ {p}")
        bad += bool(b.problems)
    lines.append("─" * 52)
    lines.append(f"{len(built)}건 처리 · 확인 필요 {bad}건")
    if bad:
        lines.append("각 job 을 check 로 한 번 더 보세요: "
                     "for f in jobs/*.json; do ./run.sh check \"$f\"; done")
    return "\n".join(lines)
